# -*- coding: utf-8 -*-
"""Helpers + document setup, phong cach giong docs/lidar_module_phan_tich.docx"""
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# <repo>/tools/docgen/common.py -> <repo>
BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, os.pardir))
HPP_PATH = os.path.join(BASE, "src", "camera_node", "camera_node.hpp")
CPP_PATH = os.path.join(BASE, "src", "camera_node", "camera_node.cpp")
FUSION_PATH = os.path.join(BASE, "src", "fusion_node", "fusion_viz_node.cpp")
OUT = os.path.join(BASE, "docs", "camera_lane_phan_tich.docx")

HPP = open(HPP_PATH, encoding="utf-8", errors="replace").read().splitlines()
CPP = open(CPP_PATH, encoding="utf-8", errors="replace").read().splitlines()
FUS = open(FUSION_PATH, encoding="utf-8", errors="replace").read().splitlines()

CONTENT_W = 9746  # twips

doc = Document()

_st = doc.styles["Normal"]
_st.font.name = "Calibri"
_st.font.size = Pt(11)
_st.paragraph_format.space_after = Pt(4)

_sec = doc.sections[0]
_sec.page_width = Emu(7772400)
_sec.page_height = Emu(10058400)
_sec.left_margin = Emu(791845)
_sec.right_margin = Emu(791845)
_sec.top_margin = Emu(720090)
_sec.bottom_margin = Emu(720090)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _add_inline(par, text, size=19, bold=False, mono=False):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _add_inline(par, part[2:-2], size=size, bold=True, mono=mono)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1] or " ")
            r.font.size = Pt(size / 2.0 - 1)
            r.font.bold = bold
            r.font.name = "Consolas"
        else:
            r = par.add_run(part if part else " ")
            r.font.size = Pt(size / 2.0 - (1 if mono else 0))
            r.font.bold = bold
            r.font.name = "Consolas" if mono else "Calibri"


def set_cell(cell, text, size=19, bold=False, mono=False, after=20):
    cell.text = ""
    par = cell.paragraphs[0]
    lines = str(text).split("\n")
    for idx, ln in enumerate(lines):
        if idx:
            par = cell.add_paragraph()
        par.paragraph_format.space_after = Pt(after / 20.0)
        par.paragraph_format.space_before = Pt(0)
        _add_inline(par, ln, size=size, bold=bold, mono=mono)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def mk_table(headers, rows, fracs=None, mono_cols=(), after=60):
    ncols = len(headers) if headers else len(rows[0])
    if fracs is None:
        fracs = [1.0 / ncols] * ncols
    total = float(sum(fracs))
    widths = [int(CONTENT_W * f / total) for f in fracs]
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if headers:
        hr = t.add_row()
        for i, h in enumerate(headers):
            set_cell(hr.cells[i], h, bold=True, after=40)
            shade(hr.cells[i], "DCE3F0")
            hr.cells[i].width = Emu(widths[i] * 635)
        repeat_header(hr)
    for r in rows:
        tr = t.add_row()
        for i, v in enumerate(r):
            set_cell(tr.cells[i], v, mono=(i in mono_cols))
            tr.cells[i].width = Emu(widths[i] * 635)
    p(after=after)
    return t


def code(lines, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(1)
        r = cp.add_run(caption)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor.from_string("1F3050")
        r.font.bold = True
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Emu(CONTENT_W * 635)
    shade(cell, "F2F4F7")
    cell.text = ""
    first = True
    for ln in lines:
        par = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        r = par.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
    p(after=80)
    return t


def note(items, fill="FFF4E5", title=None):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Emu(CONTENT_W * 635)
    shade(cell, fill)
    cell.text = ""
    allitems = []
    if title:
        allitems.append((title, True))
    if isinstance(items, str):
        items = [items]
    for it in items:
        allitems.append((it, False))
    first = True
    for txt, is_title in allitems:
        par = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        par.paragraph_format.space_after = Pt(2)
        _add_inline(par, txt, size=19, bold=is_title)
        if is_title:
            for r in par.runs:
                r.font.color.rgb = RGBColor.from_string("9C5700")
    p(after=80)
    return t


TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_runs(par, text, size=21, color=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.font.bold = True
            r.font.size = Pt(size / 2.0)
            r.font.name = "Calibri"
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(size / 2.0 - 1)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        else:
            r = par.add_run(part)
            r.font.size = Pt(size / 2.0)
            r.font.name = "Calibri"
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


def p(text="", size=21, bold=False, color=None, align=None, after=80):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(after / 20.0)
    par.paragraph_format.space_before = Pt(0)
    if align == "c":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if bold:
        add_runs(par, text, size=size, color=color)
        for r in par.runs:
            r.font.bold = True
    else:
        add_runs(par, text, size=size, color=color)
    return par


def h1(text):
    return doc.add_heading(text, level=1)


def h2(text):
    return doc.add_heading(text, level=2)


def h3(text):
    return doc.add_heading(text, level=3)


def bullets(items, style="List Bullet"):
    for it in items:
        par = doc.add_paragraph(style=style)
        par.paragraph_format.space_after = Pt(2)
        add_runs(par, it)


def pagebreak():
    doc.add_page_break()


def rng(lines, a, b, numbered=True, prefix=None):
    out = []
    for i in range(a, b + 1):
        txt = lines[i - 1]
        if numbered:
            out.append("%4d  %s" % (i, txt))
        else:
            out.append(txt)
    return out
